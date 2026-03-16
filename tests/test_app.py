from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app import app, db, init_db, Student, TeacherAccount, SignInRecord
from werkzeug.security import check_password_hash


def login_teacher(client, password="abc123"):
    return client.post("/teacher/login", data={"password": password}, follow_redirects=False)


def setup_module():
    app.config["TESTING"] = True
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///:memory:"
    with app.app_context():
        db.drop_all()
        db.create_all()
        init_db()
        db.session.add(Student(name="张三"))
        db.session.add(Student(name="李四"))
        db.session.commit()


def test_teacher_login_and_password_change():
    client = app.test_client()

    no_auth = client.get("/api/status")
    assert no_auth.status_code == 401

    bad = login_teacher(client, "bad")
    assert bad.status_code == 401

    ok = login_teacher(client)
    assert ok.status_code == 302

    change = client.post("/api/change-password", json={"old_password": "abc123", "new_password": "newpass1"})
    assert change.status_code == 200
    assert change.json["ok"] is True

    client.post("/teacher/logout")
    relogin = login_teacher(client, "newpass1")
    assert relogin.status_code == 302

    with app.app_context():
        account = db.session.get(TeacherAccount, 1)
        assert account is not None
        assert check_password_hash(account.password_hash, "newpass1")


def test_student_api_and_signin():
    client = app.test_client()

    res = client.get("/api/students?q=张")
    assert res.status_code == 200
    assert "张三" in res.json["students"]

    signin = client.post(
        "/api/signin",
        json={"student_name": "张三"},
        environ_base={"REMOTE_ADDR": "10.0.0.20"},
    )
    assert signin.status_code == 200
    assert signin.json["ok"] is True

    duplicate = client.post("/api/signin", json={"student_name": "张三"})
    assert duplicate.status_code == 400
    assert duplicate.json["ok"] is False

    with app.app_context():
        record = SignInRecord.query.filter_by(student_name="张三").first()
        assert record is not None
        assert record.computer_name == "10.0.0.20"


def test_import_export_and_status():
    client = app.test_client()
    login_teacher(client, "newpass1")

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["王五"])
    buff = BytesIO()
    wb.save(buff)
    buff.seek(0)

    import_res = client.post(
        "/api/import",
        data={"file": (buff, "students.xlsx")},
        content_type="multipart/form-data",
    )
    assert import_res.status_code == 200


    with app.app_context():
        if not SignInRecord.query.filter_by(student_name="李四").first():
            db.session.add(
                SignInRecord(
                    student_name="李四",
                    computer_name="10.0.0.30",
                    signed_at=datetime.now() + timedelta(seconds=1),
                )
            )
            db.session.commit()

    status = client.get("/api/status")
    assert status.status_code == 200
    assert status.json["total_count"] >= 3
    assert any(row["student_name"] == "王五" and row["status"] == "未签到" for row in status.json["roster"])
    signed_rows = [row for row in status.json["roster"] if row["status"] == "已签到"]
    assert len(signed_rows) >= 2
    assert signed_rows[0]["student_name"] == "张三"
    assert signed_rows[1]["student_name"] == "李四"

    export = client.get("/api/export")
    assert export.status_code == 200
    assert export.headers["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(filename=BytesIO(export.data))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(row[0] == "王五" and row[1] == "未签到" and row[3] in (None, "") for row in rows)


def test_clear_roster():
    client = app.test_client()
    login_teacher(client, "newpass1")

    clear_res = client.post("/api/clear-roster")
    assert clear_res.status_code == 200
    assert clear_res.json["ok"] is True

    status = client.get("/api/status")
    assert status.status_code == 200
    assert status.json["total_count"] == 0
    assert status.json["signed_count"] == 0
    assert status.json["roster"] == []


def test_seat_map_page():
    client = app.test_client()
    login_teacher(client, "newpass1")

    with app.app_context():
        if not SignInRecord.query.filter_by(student_name="座位测试A").first():
            db.session.add(SignInRecord(student_name="座位测试A", computer_name="10.10.10.60", signed_at=datetime.now()))
        if not SignInRecord.query.filter_by(student_name="座位测试B").first():
            db.session.add(SignInRecord(student_name="座位测试B", computer_name="10.10.10.15", signed_at=datetime.now()))
        db.session.commit()

    page = client.get('/teacher/seats')
    assert page.status_code == 200
    text = page.get_data(as_text=True)
    assert '#60' in text
    assert '#15' in text
    assert '座位测试A' in text
    assert '座位测试B' in text
